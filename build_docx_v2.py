import io, os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK

SRC = r"D:\empty\KDP-POLISHED"
FILES = [f for f in sorted(os.listdir(SRC)) if re.match(r'^0[1-7]_.*\.txt$', f)]

def base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(11.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25

def add_pagebreak(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def para(doc, text, align=None, size=None, bold=False, italic=False, before=0, after=6):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def heading(doc, text, level):
    h = doc.add_heading("", level=level)
    h.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    h.paragraph_format.space_after = Pt(8)
    r = h.add_run(text)
    r.font.name = "Georgia"
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)

def code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(10)

def build(txt_path, out_path):
    lines = io.open(txt_path, encoding="utf-8").read().splitlines()
    n = len(lines)
    # pass 1: collect headings for Contents
    h1s, h2s = [], []
    i = 0
    while i < n:
        s = lines[i].strip()
        if s.startswith("# ") and not s.startswith("## "): h1s.append(s[2:].strip())
        elif s.startswith("## "): h2s.append(s[3:].strip())
        i += 1
    # pass 2: emit
    doc = Document()
    base_style(doc)
    cp = doc.core_properties
    cp.author = "Niloy Das"
    i = 0
    title_done = copy_done = contents_done = False
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if s == "@TITLE":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDTITLE":
                if lines[i].strip(): block.append(lines[i].strip())
                i += 1
            i += 1
            for k, t in enumerate(block):
                if k == 0:    para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=30, bold=True, before=180, after=14)
                elif k == 1:  para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=15, italic=True, after=200)
                elif k == 2:  para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, after=8)
                else:         para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=12, after=6)
            cp.title = block[0]
            add_pagebreak(doc); title_done = True
            continue
        if s == "@COPY":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDCOPY":
                if lines[i].strip(): block.append(lines[i].strip())
                i += 1
            i += 1
            for t in block:
                para(doc, t, size=10, after=3)
            add_pagebreak(doc); copy_done = True
            # static Contents
            if h1s:
                heading(doc, "Contents", 1)
                h2set = set(h2s)
                used_h2 = False
                for h in h1s:
                    if h == "Contents": continue
                    para(doc, h, size=11.5, bold=True, before=8, after=2)
                add_pagebreak(doc)
            contents_done = True
            continue
        if s == "@PAGEBREAK":
            add_pagebreak(doc); i += 1; continue
        if s == "@CODE":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDCODE":
                block.append(lines[i].rstrip()); i += 1
            i += 1
            code_block(doc, [b for b in block if b])
            continue
        if s == "@CENTER":
            i += 1
            while i < n and lines[i].strip() == "@CENTER": i += 1
            block = []
            while i < n and lines[i].strip() and not lines[i].strip().startswith("@"):
                block.append(lines[i].strip()); i += 1
            for t in block:
                para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, before=14, after=4)
            continue
        if s.startswith("@CENTER "):
            para(doc, s[8:].strip(), WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, before=14, after=4)
            i += 1; continue
        if s.startswith("# ") and not s.startswith("## "):
            heading(doc, s[2:].strip(), 1); i += 1; continue
        if s.startswith("## "):
            heading(doc, s[3:].strip(), 2); i += 1; continue
        if s.startswith("### "):
            heading(doc, s[4:].strip(), 3); i += 1; continue
        if s.startswith("> "):
            para(doc, s[2:].strip(), WD_ALIGN_PARAGRAPH.CENTER, italic=True, before=10, after=10)
            i += 1; continue
        if s == "---":
            para(doc, "\u2726  \u2726  \u2726", WD_ALIGN_PARAGRAPH.CENTER, before=12, after=12)
            i += 1; continue
        para(doc, s)
        i += 1
    doc.save(out_path)
    return len(h1s)

for f in FILES:
    src = os.path.join(SRC, f)
    out = os.path.join(SRC, os.path.splitext(f)[0] + ".docx")
    n = build(src, out)
    print("built", os.path.basename(out), os.path.getsize(out), "bytes,", n, "H1s")
