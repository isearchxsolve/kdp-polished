import io, os, re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = r"D:\empty\KDP-POLISHED"
FILES = [f for f in sorted(os.listdir(SRC)) if re.match(r"^0[1-7]_.*\.txt$", f)]

def base_style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(11.5)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.line_spacing = 1.25

def add_pagebreak(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def para(doc, text, align=None, size=None, bold=False, italic=False, before=0, after=6):
    p = doc.add_paragraph()
    if align: p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text)
    r.bold = bold; r.italic = italic
    if size: r.font.size = Pt(size)
    return p

def add_bookmark(paragraph, name, bid):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid)); start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)

def add_toc_link(doc, text, anchor, level=1, size=11.5, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Pt(24 if level == 2 else 0)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "1A1A2E")
    rPr.append(color)
    if bold:
        b = OxmlElement("w:b"); rPr.append(b)
    sz = OxmlElement("w:sz"); sz.set(qn("w:val"), str(int(size * 2))); rPr.append(sz)
    fonts = OxmlElement("w:rFonts"); fonts.set(qn("w:ascii"), "Georgia"); fonts.set(qn("w:hAnsi"), "Georgia")
    rPr.append(fonts)
    r.append(rPr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    r.append(t)
    hl.append(r)
    p._p.append(hl)

def toc_field_start(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    fld.set(qn("w:dirty"), "true")
    r._r.append(fld)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    r._r.append(instr)
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r._r.append(sep)

def toc_field_end(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r._r.append(end)

def set_update_fields(doc):
    el = doc.settings.element
    uf = OxmlElement("w:updateFields")
    uf.set(qn("w:val"), "true")
    el.append(uf)

def heading(doc, text, level, bid):
    h = doc.add_heading("", level=level)
    h.paragraph_format.space_before = Pt(18 if level == 1 else 12)
    h.paragraph_format.space_after = Pt(8)
    r = h.add_run(text)
    r.font.name = "Georgia"
    r.font.color.rgb = RGBColor(0x1A, 0x1A, 0x2E)
    add_bookmark(h, f"toc_{bid}", bid)

def code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(ln)
        r.font.name = "Consolas"
        r.font.size = Pt(10)

def build(txt_path, out_path):
    lines = io.open(txt_path, encoding="utf-8-sig").read().splitlines()
    n = len(lines)
    # pass 1: ordered heading list (H1/H2) for bookmarks + contents
    toc = []
    for l in lines:
        s = l.strip()
        if s.startswith("## "):
            toc.append((2, s[3:].strip()))
        elif s.startswith("# ") and not s.startswith("## "):
            toc.append((1, s[2:].strip()))
    doc = Document()
    base_style(doc)
    cp = doc.core_properties
    cp.author = "Niloy Das"
    i = 0
    next_bid = 100
    title_done = False
    W_front = True
    while i < n:
        s = lines[i].strip()
        if not s:
            i += 1; continue
        if s == "@TITLE":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDTITLE":
                if lines[i].strip(): block.append(lines[i].strip())
                i += 1
            i += 1
            for k, t in enumerate(block):
                if k == 0:
                    para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=30, bold=True, before=180, after=14)
                elif k == 1:
                    para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=15, italic=True, after=200)
                elif k == 2:
                    para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, after=8)
                else:
                    para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=12, after=6)
            cp.title = block[0]
            add_pagebreak(doc); title_done = True
            continue
        if s == "@COPY":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDCOPY":
                if lines[i].strip(): block.append(lines[i].strip())
                i += 1
            i += 1
            for t in block:
                para(doc, t, size=10, after=3)
            add_pagebreak(doc)
            # Contents page: Word TOC field wrapping linked entries (KDP-recognizable)
            if toc:
                heading(doc, "Contents", 1, next_bid); next_bid += 1
                toc_field_start(doc)
                for idx, (lv, h) in enumerate(toc):
                    if h == "Contents": continue
                    add_toc_link(doc, h, f"toc_{idx + 101}", level=lv, size=11.5 if lv == 1 else 10.5, bold=(lv == 1))
                toc_field_end(doc)
                set_update_fields(doc)
                add_pagebreak(doc)
            continue
        if s == "@PAGEBREAK":
            add_pagebreak(doc); i += 1; continue
        if s == "@CODE":
            block = []
            i += 1
            while i < n and lines[i].strip() != "@ENDCODE":
                block.append(lines[i].rstrip()); i += 1
            i += 1
            code_block(doc, [b for b in block if b])
            continue
        if s == "@CENTER":
            i += 1
            while i < n and lines[i].strip() == "@CENTER": i += 1
            block = []
            while i < n and lines[i].strip() and not lines[i].strip().startswith("@"):
                block.append(lines[i].strip()); i += 1
            for t in block:
                para(doc, t, WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, before=14, after=4)
            continue
        if s.startswith("@CENTER "):
            para(doc, s[8:].strip(), WD_ALIGN_PARAGRAPH.CENTER, size=12, italic=True, before=14, after=4)
            i += 1; continue
        if s.startswith("# ") and not s.startswith("## "):
            heading(doc, s[2:].strip(), 1, next_bid); next_bid += 1
            i += 1; continue
        if s.startswith("## "):
            heading(doc, s[3:].strip(), 2, next_bid); next_bid += 1
            i += 1; continue
        if s.startswith("### "):
            heading(doc, s[4:].strip(), 3, next_bid); next_bid += 1
            i += 1; continue
        if s.startswith("> "):
            para(doc, s[2:].strip(), WD_ALIGN_PARAGRAPH.CENTER, italic=True, before=10, after=10)
            i += 1; continue
        if s == "---":
            para(doc, "\u2726  \u2726  \u2726", WD_ALIGN_PARAGRAPH.CENTER, before=12, after=12)
            i += 1; continue
        para(doc, s)
        i += 1
    doc.save(out_path)
    return len(toc)

for f in FILES:
    src = os.path.join(SRC, f)
    out = os.path.join(SRC, os.path.splitext(f)[0] + ".docx")
    n = build(src, out)
    print("built", os.path.basename(out), os.path.getsize(out), "bytes,", n, "TOC entries")
